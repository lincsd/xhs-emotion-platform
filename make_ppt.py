#!/usr/bin/env python3
"""生成「小红书智能运营台」产品介绍 PPT"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# ── 配色 ──
CLR_BG       = RGBColor(0x0F, 0x0F, 0x0F)
CLR_BG2      = RGBColor(0x1A, 0x1A, 0x2E)
CLR_ACCENT   = RGBColor(0xFF, 0x2E, 0x63)  # 小红书红
CLR_ACCENT2  = RGBColor(0xFF, 0x6B, 0x6B)
CLR_GOLD     = RGBColor(0xFF, 0xD7, 0x00)
CLR_GREEN    = RGBColor(0x2E, 0xCC, 0x71)
CLR_BLUE     = RGBColor(0x36, 0x9B, 0xFF)
CLR_PURPLE   = RGBColor(0xA8, 0x5C, 0xFF)
CLR_WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
CLR_GRAY     = RGBColor(0xBB, 0xBB, 0xBB)
CLR_LIGHT    = RGBColor(0xF0, 0xF0, 0xF0)
CLR_DARK_CARD= RGBColor(0x1E, 0x1E, 0x3A)
CLR_CARD_BG  = RGBColor(0x25, 0x25, 0x45)

def set_slide_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_shape_bg(slide, left, top, width, height, color, radius=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    if radius is not None:
        shape.adjustments[0] = radius
    return shape

def add_textbox(slide, left, top, width, height, text, font_size=18,
                color=CLR_WHITE, bold=False, align=PP_ALIGN.LEFT, font_name='微软雅黑'):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = align
    return txBox

def add_multiline(slide, left, top, width, height, lines, font_size=16,
                  color=CLR_WHITE, bold=False, line_spacing=1.5, align=PP_ALIGN.LEFT):
    """lines: list of (text, color, bold, font_size) or just str"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(lines):
        if isinstance(item, str):
            txt, clr, b, fs = item, color, bold, font_size
        else:
            txt = item[0]
            clr = item[1] if len(item) > 1 else color
            b = item[2] if len(item) > 2 else bold
            fs = item[3] if len(item) > 3 else font_size
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = txt
        p.font.size = Pt(fs)
        p.font.color.rgb = clr
        p.font.bold = b
        p.font.name = '微软雅黑'
        p.alignment = align
        p.space_after = Pt(font_size * (line_spacing - 1))
    return txBox

def add_card(slide, left, top, width, height, icon, title, desc, accent=CLR_ACCENT):
    card = add_shape_bg(slide, left, top, width, height, CLR_CARD_BG, 0.05)
    # icon
    add_textbox(slide, left + Inches(0.3), top + Inches(0.25), Inches(0.8), Inches(0.8),
                icon, font_size=36, align=PP_ALIGN.LEFT)
    # title
    add_textbox(slide, left + Inches(0.3), top + Inches(0.85), width - Inches(0.6), Inches(0.5),
                title, font_size=18, color=accent, bold=True)
    # desc
    add_multiline(slide, left + Inches(0.3), top + Inches(1.3), width - Inches(0.6), height - Inches(1.5),
                  desc if isinstance(desc, list) else [desc], font_size=13, color=CLR_GRAY, line_spacing=1.4)
    return card


# ═══════════════════════════════════════
# SLIDE 1 — 封面
# ═══════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s, CLR_BG)

# 装饰圆
circle = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.5), Inches(-1), Inches(5), Inches(5))
circle.fill.solid(); circle.fill.fore_color.rgb = RGBColor(0xFF, 0x2E, 0x63)
circle.line.fill.background()
import copy
from lxml import etree
NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'
def set_shape_alpha(shape, pct):
    """Set fill alpha (0-100) on a solid-filled shape."""
    spPr = shape._element.find(f'.//{{{NS}}}solidFill')
    if spPr is not None:
        srgb = spPr.find(f'{{{NS}}}srgbClr')
        if srgb is not None:
            a = etree.SubElement(srgb, f'{{{NS}}}alpha')
            a.set('val', str(pct * 1000))
set_shape_alpha(circle, 15)

circle2 = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(-2), Inches(4), Inches(6), Inches(6))
circle2.fill.solid(); circle2.fill.fore_color.rgb = CLR_PURPLE
circle2.line.fill.background()
set_shape_alpha(circle2, 10)

add_textbox(s, Inches(1.5), Inches(1.2), Inches(10), Inches(1),
            '📝 小红书智能运营台', font_size=52, color=CLR_WHITE, bold=True, align=PP_ALIGN.LEFT)

add_textbox(s, Inches(1.5), Inches(2.5), Inches(9), Inches(0.8),
            'AI 驱动的小红书多技能内容创作与运营管理平台', font_size=26, color=CLR_ACCENT2, align=PP_ALIGN.LEFT)

add_multiline(s, Inches(1.5), Inches(3.8), Inches(9), Inches(2.5), [
    ('✦  7 大账号技能模板    ✦  Gemini AI 智能生成    ✦  一站式内容管理', CLR_GRAY, False, 18),
    ('✦  A/B 测试优化        ✦  数据看板分析          ✦  SEO 标题/标签生成', CLR_GRAY, False, 18),
    ('✦  品牌定位分析        ✦  竞品搜索分析          ✦  AI 配图自动生成', CLR_GRAY, False, 18),
], font_size=18, color=CLR_GRAY, line_spacing=2.0)

# 底部信息
add_textbox(s, Inches(1.5), Inches(6.5), Inches(6), Inches(0.5),
            '产品介绍  ·  2026', font_size=14, color=RGBColor(0x66, 0x66, 0x66))


# ═══════════════════════════════════════
# SLIDE 2 — 产品定位：这是什么？
# ═══════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s, CLR_BG)

add_textbox(s, Inches(0.8), Inches(0.4), Inches(6), Inches(0.8),
            '🎯 产品定位', font_size=36, color=CLR_WHITE, bold=True)
add_shape_bg(s, Inches(0.8), Inches(1.1), Inches(2), Inches(0.06), CLR_ACCENT)

add_shape_bg(s, Inches(0.8), Inches(1.6), Inches(11.5), Inches(2.8), CLR_DARK_CARD, 0.03)
add_multiline(s, Inches(1.3), Inches(1.8), Inches(10.5), Inches(2.5), [
    ('「小红书智能运营台」是一款面向小红书创作者、运营人员和自媒体从业者的', CLR_WHITE, False, 20),
    ('  AI 驱动的一站式内容创作与运营管理平台。', CLR_ACCENT, True, 20),
    ('', CLR_WHITE, False, 10),
    ('它利用 Google Gemini AI 大模型，结合小红书平台特点和热门趋势，帮助用户快速生成', CLR_GRAY, False, 17),
    ('高质量、多风格的种草笔记，并提供从内容策划到数据分析的全链路管理。', CLR_GRAY, False, 17),
], line_spacing=1.6)

# 3大价值卡片
vals = [
    ('🚀', '效率提升', '用 AI 替代繁琐的选题、写作流程\n内容生成速度提升 10 倍', CLR_ACCENT),
    ('🎨', '多元技能', '覆盖 7 大赛道技能模板\n一个平台管理多个账号定位', CLR_PURPLE),
    ('📊', '数据驱动', '搜索分析 + A/B 测试 + 看板\n用数据指导每一篇内容决策', CLR_BLUE),
]
for i, (icon, title, desc, clr) in enumerate(vals):
    x = Inches(0.8 + i * 4.0)
    add_card(s, x, Inches(4.8), Inches(3.6), Inches(2.4), icon, title, desc.split('\n'), clr)


# ═══════════════════════════════════════
# SLIDE 3 — 核心功能总览
# ═══════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s, CLR_BG)

add_textbox(s, Inches(0.8), Inches(0.4), Inches(8), Inches(0.8),
            '⚡ 核心功能总览', font_size=36, color=CLR_WHITE, bold=True)
add_shape_bg(s, Inches(0.8), Inches(1.1), Inches(2), Inches(0.06), CLR_ACCENT)

features = [
    ('🤖', 'AI 内容生成', ['Gemini AI 一键生成完整笔记', '标题 / 正文 / 标签全自动', '支持参考笔记学习热门风格'], CLR_ACCENT),
    ('🔍', '搜索分析引擎', ['Google + 百度双引擎搜索', '自动分析小红书热门内容', '提炼爆款规律和趋势'], CLR_BLUE),
    ('🧪', 'A/B 测试', ['同主题生成多版本标题/内容', '对比选出最优方案', '数据化提升点击率'], CLR_GREEN),
    ('🖼️', 'AI 配图生成', ['Gemini 图像生成能力', '自动生成笔记封面/内容图', '省去寻找配图的时间'], CLR_PURPLE),
    ('📊', '数据看板', ['收入追踪与趋势分析', '笔记发布统计', '可视化图表一目了然'], CLR_GOLD),
    ('🏷️', '品牌定位分析', ['AI 分析账号品牌定位', '竞品对标 & 差异化建议', '输出可执行的运营策略'], RGBColor(0xFF, 0x85, 0x00)),
]
for i, (icon, title, descs, clr) in enumerate(features):
    col = i % 3
    row = i // 3
    x = Inches(0.6 + col * 4.1)
    y = Inches(1.6 + row * 3.0)
    add_card(s, x, y, Inches(3.8), Inches(2.6), icon, title, descs, clr)


# ═══════════════════════════════════════
# SLIDE 4 — 7 大技能模板
# ═══════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s, CLR_BG)

add_textbox(s, Inches(0.8), Inches(0.3), Inches(10), Inches(0.8),
            '🎭 7 大技能模板 — 覆盖热门赛道', font_size=36, color=CLR_WHITE, bold=True)
add_shape_bg(s, Inches(0.8), Inches(1.0), Inches(2), Inches(0.06), CLR_ACCENT)

skills = [
    ('💕', '情感', '情感故事/恋爱心理\n治愈系/成长感悟', RGBColor(0xFF, 0x69, 0xB4)),
    ('🍜', '美食', '探店/食谱/测评\n美食攻略/种草', RGBColor(0xFF, 0x8C, 0x00)),
    ('✈️', '旅行', '旅行攻略/打卡分享\n小众景点/行程规划', CLR_BLUE),
    ('💪', '健身', '健身教程/饮食计划\n身材对比/运动打卡', CLR_GREEN),
    ('👗', '穿搭', 'OOTD/风格搭配\n季节穿搭/好物推荐', CLR_PURPLE),
    ('📱', '科技数码', '数码测评/好物开箱\n使用技巧/对比分析', CLR_BLUE),
    ('🛒', '电商带货', '美妆/家居/母婴/数码\n零食/服饰/个护/白菜', CLR_ACCENT),
]

for i, (icon, name, desc, clr) in enumerate(skills):
    if i < 4:
        x = Inches(0.5 + i * 3.1)
        y = Inches(1.5)
    else:
        x = Inches(2.0 + (i - 4) * 3.1)
        y = Inches(4.2)

    card = add_shape_bg(s, x, y, Inches(2.8), Inches(2.3), CLR_CARD_BG, 0.05)
    add_textbox(s, x + Inches(0.2), y + Inches(0.2), Inches(1), Inches(0.8),
                icon, font_size=40, align=PP_ALIGN.LEFT)
    add_textbox(s, x + Inches(0.9), y + Inches(0.25), Inches(1.5), Inches(0.5),
                name, font_size=22, color=clr, bold=True)
    add_multiline(s, x + Inches(0.2), y + Inches(1.0), Inches(2.4), Inches(1.2),
                  desc.split('\n'), font_size=13, color=CLR_GRAY, line_spacing=1.5)

add_textbox(s, Inches(0.8), Inches(6.8), Inches(11), Inches(0.5),
            '每个技能内置独立的内容模板、分类体系和AI提示词，确保生成内容贴合赛道特点',
            font_size=14, color=RGBColor(0x88, 0x88, 0x88), align=PP_ALIGN.LEFT)


# ═══════════════════════════════════════
# SLIDE 5 — AI 工作流演示
# ═══════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s, CLR_BG)

add_textbox(s, Inches(0.8), Inches(0.4), Inches(10), Inches(0.8),
            '🔄 AI 内容生成工作流', font_size=36, color=CLR_WHITE, bold=True)
add_shape_bg(s, Inches(0.8), Inches(1.1), Inches(2), Inches(0.06), CLR_ACCENT)

steps = [
    ('① 输入主题', '用户输入关键词\n如"秋冬护肤"', CLR_ACCENT),
    ('② 搜索分析', 'Google + 百度双引擎\n分析小红书热门内容', CLR_BLUE),
    ('③ AI 创作', 'Gemini AI 生成\n标题 + 正文 + 标签', CLR_PURPLE),
    ('④ 配图生成', 'AI 自动生成\n匹配内容的精美配图', CLR_GREEN),
    ('⑤ 优化发布', 'A/B 测试对比\nSEO 优化后发布', CLR_GOLD),
]

for i, (title, desc, clr) in enumerate(steps):
    x = Inches(0.5 + i * 2.5)
    y = Inches(1.8)

    card = add_shape_bg(s, x, y, Inches(2.2), Inches(2.5), CLR_CARD_BG, 0.05)
    add_textbox(s, x + Inches(0.15), y + Inches(0.2), Inches(1.9), Inches(0.5),
                title, font_size=18, color=clr, bold=True, align=PP_ALIGN.CENTER)
    add_multiline(s, x + Inches(0.15), y + Inches(0.9), Inches(1.9), Inches(1.5),
                  desc.split('\n'), font_size=14, color=CLR_GRAY, line_spacing=1.6, align=PP_ALIGN.CENTER)

    # 箭头 (除了最后一个)
    if i < len(steps) - 1:
        add_textbox(s, x + Inches(2.15), y + Inches(0.9), Inches(0.5), Inches(0.5),
                    '→', font_size=28, color=RGBColor(0x55, 0x55, 0x55), align=PP_ALIGN.CENTER)

# 底部：参考笔记功能说明
add_shape_bg(s, Inches(0.8), Inches(4.8), Inches(11.5), Inches(2.2), CLR_DARK_CARD, 0.03)
add_textbox(s, Inches(1.2), Inches(4.95), Inches(5), Inches(0.5),
            '📌 独特亮点：参考笔记 + 双搜索引擎', font_size=20, color=CLR_GOLD, bold=True)
add_multiline(s, Inches(1.2), Inches(5.5), Inches(10.5), Inches(1.5), [
    ('▸ 用户可粘贴真实爆款笔记作为参考，AI 会学习其风格和结构进行创作', CLR_GRAY, False, 15),
    ('▸ Google + 百度双搜索引擎采集 site:xiaohongshu.com 实时热门内容', CLR_GRAY, False, 15),
    ('▸ 三层信息叠加：搜索趋势 + 参考笔记 + 用户需求 = 高质量输出', CLR_GRAY, False, 15),
], line_spacing=1.6)


# ═══════════════════════════════════════
# SLIDE 6 — 适合谁用？
# ═══════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s, CLR_BG)

add_textbox(s, Inches(0.8), Inches(0.4), Inches(10), Inches(0.8),
            '👥 适合谁用？', font_size=36, color=CLR_WHITE, bold=True)
add_shape_bg(s, Inches(0.8), Inches(1.1), Inches(2), Inches(0.06), CLR_ACCENT)

personas = [
    ('🌟', '小红书博主 / KOL', '需要持续产出高质量内容的创作者\n每天更新多篇笔记，需要效率工具', [
        '内容创作效率提升 10x',
        '多赛道切换零成本',
        'A/B 测试提升爆文率',
    ], CLR_ACCENT),
    ('💼', '品牌运营 / MCN', '管理多个账号的运营团队\n需要标准化内容生产流程', [
        '品牌定位一键分析',
        '多技能模板统一管理',
        'SEO 优化自动完成',
    ], CLR_BLUE),
    ('🏠', '个体商家 / 电商卖家', '想在小红书做种草营销的商家\n没有专业写手但需要内容推广', [
        '电商带货专属模板',
        '零写作基础也能出稿',
        '配图自动生成省时省力',
    ], CLR_GREEN),
    ('📚', '自媒体新手 / 副业', '刚入行的新人或想做副业的人\n不懂平台调性和写作技巧', [
        '爆款分析快速学习',
        '参考笔记模仿学习',
        '每日 3 次免费 AI 使用',
    ], CLR_PURPLE),
]

for i, (icon, title, desc, points, clr) in enumerate(personas):
    col = i % 2
    row = i // 2
    x = Inches(0.6 + col * 6.2)
    y = Inches(1.5 + row * 2.9)
    w = Inches(5.8)
    h = Inches(2.6)

    card = add_shape_bg(s, x, y, w, h, CLR_CARD_BG, 0.04)
    add_textbox(s, x + Inches(0.3), y + Inches(0.2), Inches(0.8), Inches(0.7),
                icon, font_size=32)
    add_textbox(s, x + Inches(1.0), y + Inches(0.2), Inches(4), Inches(0.5),
                title, font_size=20, color=clr, bold=True)
    add_multiline(s, x + Inches(1.0), y + Inches(0.7), Inches(4.5), Inches(0.8),
                  desc.split('\n'), font_size=12, color=CLR_GRAY, line_spacing=1.4)

    for j, pt in enumerate(points):
        add_textbox(s, x + Inches(1.0), y + Inches(1.45 + j * 0.35), Inches(4.5), Inches(0.4),
                    f'✓ {pt}', font_size=13, color=CLR_WHITE)


# ═══════════════════════════════════════
# SLIDE 7 — 使用场景
# ═══════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s, CLR_BG)

add_textbox(s, Inches(0.8), Inches(0.4), Inches(10), Inches(0.8),
            '💡 典型使用场景', font_size=36, color=CLR_WHITE, bold=True)
add_shape_bg(s, Inches(0.8), Inches(1.1), Inches(2), Inches(0.06), CLR_ACCENT)

scenes = [
    ('场景 1：日常内容更新', CLR_ACCENT, [
        ('早起选题 → 输入关键词 → AI 一键生成 3 篇备选', CLR_WHITE),
        ('挑选最佳版本 → 修改润色 → 发布', CLR_GRAY),
        ('⏱ 原来 2 小时的工作，现在 15 分钟搞定', CLR_GOLD),
    ]),
    ('场景 2：爆款内容复刻', CLR_BLUE, [
        ('看到竞品爆文 → 粘贴到「参考笔记」', CLR_WHITE),
        ('AI 分析爆款结构 → 生成同风格差异化内容', CLR_GRAY),
        ('⏱ 站在巨人肩膀上，快速产出互动率高的笔记', CLR_GOLD),
    ]),
    ('场景 3：电商新品种草', CLR_GREEN, [
        ('切换「电商带货」技能 → 选择品类', CLR_WHITE),
        ('输入产品卖点 → AI 生成种草文案 + 配图', CLR_GRAY),
        ('⏱ 1 个产品 5 分钟出一篇专业种草笔记', CLR_GOLD),
    ]),
    ('场景 4：品牌矩阵运营', CLR_PURPLE, [
        ('品牌定位分析 → 确定差异化方向', CLR_WHITE),
        ('7 大技能切换管理不同定位账号', CLR_GRAY),
        ('⏱ 一个人管理多个账号也能游刃有余', CLR_GOLD),
    ]),
]

for i, (title, clr, items) in enumerate(scenes):
    col = i % 2
    row = i // 2
    x = Inches(0.6 + col * 6.2)
    y = Inches(1.5 + row * 2.8)

    card = add_shape_bg(s, x, y, Inches(5.8), Inches(2.5), CLR_CARD_BG, 0.04)
    add_textbox(s, x + Inches(0.3), y + Inches(0.2), Inches(5.2), Inches(0.5),
                title, font_size=19, color=clr, bold=True)
    add_multiline(s, x + Inches(0.3), y + Inches(0.8), Inches(5.2), Inches(1.6),
                  items, font_size=14, line_spacing=1.7)


# ═══════════════════════════════════════
# SLIDE 8 — 积分体系
# ═══════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s, CLR_BG)

add_textbox(s, Inches(0.8), Inches(0.4), Inches(10), Inches(0.8),
            '💰 灵活的积分计费体系', font_size=36, color=CLR_WHITE, bold=True)
add_shape_bg(s, Inches(0.8), Inches(1.1), Inches(2), Inches(0.06), CLR_ACCENT)

# 免费额度说明
add_shape_bg(s, Inches(0.8), Inches(1.5), Inches(5.5), Inches(1.5), CLR_DARK_CARD, 0.03)
add_textbox(s, Inches(1.2), Inches(1.6), Inches(4.5), Inches(0.5),
            '🎁 免费额度', font_size=22, color=CLR_GREEN, bold=True)
add_multiline(s, Inches(1.2), Inches(2.15), Inches(4.5), Inches(0.8), [
    ('每日 3 次免费 AI 调用，零成本体验', CLR_WHITE, False, 16),
    ('支持兑换码充值积分', CLR_GRAY, False, 14),
], line_spacing=1.5)

# 积分表
costs = [
    ('内容搜索分析', '2 积分', CLR_BLUE),
    ('AI 内容创作', '2 积分', CLR_PURPLE),
    ('图片生成', '3 积分/张', CLR_GREEN),
    ('标题优化', '1 积分', CLR_GRAY),
    ('内容改写', '1 积分', CLR_GRAY),
    ('A/B 测试', '2 积分', CLR_GOLD),
    ('品牌定位', '2 积分', RGBColor(0xFF, 0x85, 0x00)),
    ('SEO 标签生成', '1 积分', CLR_BLUE),
]

add_shape_bg(s, Inches(6.8), Inches(1.5), Inches(5.8), Inches(5.5), CLR_DARK_CARD, 0.03)
add_textbox(s, Inches(7.2), Inches(1.6), Inches(5), Inches(0.5),
            '📋 功能积分价目表', font_size=20, color=CLR_ACCENT, bold=True)

for i, (name, cost, clr) in enumerate(costs):
    y_pos = Inches(2.25 + i * 0.55)
    # 分隔条
    if i > 0:
        add_shape_bg(s, Inches(7.2), y_pos - Inches(0.05), Inches(5), Inches(0.01), RGBColor(0x33, 0x33, 0x55))
    add_textbox(s, Inches(7.2), y_pos, Inches(3), Inches(0.4),
                name, font_size=15, color=CLR_WHITE)
    add_textbox(s, Inches(10.5), y_pos, Inches(1.8), Inches(0.4),
                cost, font_size=15, color=clr, bold=True, align=PP_ALIGN.RIGHT)

# 一键生成说明
add_shape_bg(s, Inches(0.8), Inches(3.4), Inches(5.5), Inches(2.0), CLR_DARK_CARD, 0.03)
add_textbox(s, Inches(1.2), Inches(3.5), Inches(4.5), Inches(0.5),
            '📝 一键生成费用说明', font_size=18, color=CLR_ACCENT2, bold=True)
add_multiline(s, Inches(1.2), Inches(4.1), Inches(4.8), Inches(1.2), [
    ('一键生成 = 搜索分析(2) + 内容创作(2)', CLR_WHITE, False, 15),
    ('≈ 4 积分/篇（不含配图）', CLR_GOLD, True, 15),
    ('含配图：额外 +3 积分/张', CLR_GRAY, False, 14),
], line_spacing=1.6)

# 安全保障
add_shape_bg(s, Inches(0.8), Inches(5.8), Inches(5.5), Inches(1.2), CLR_DARK_CARD, 0.03)
add_textbox(s, Inches(1.2), Inches(5.9), Inches(4.5), Inches(0.5),
            '🔒 安全保障', font_size=18, color=CLR_GREEN, bold=True)
add_multiline(s, Inches(1.2), Inches(6.4), Inches(4.8), Inches(0.6), [
    ('手机号实名注册 · 数据加密存储 · 多端同步', CLR_GRAY, False, 14),
], line_spacing=1.4)


# ═══════════════════════════════════════
# SLIDE 9 — 技术优势
# ═══════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s, CLR_BG)

add_textbox(s, Inches(0.8), Inches(0.4), Inches(10), Inches(0.8),
            '🛠️ 技术优势 & 产品亮点', font_size=36, color=CLR_WHITE, bold=True)
add_shape_bg(s, Inches(0.8), Inches(1.1), Inches(2), Inches(0.06), CLR_ACCENT)

techs = [
    ('🧠', 'Gemini 2.5 Flash', '采用 Google 最新 Gemini 2.5 Flash\nThinking 模型，推理更深更准', CLR_ACCENT),
    ('🌐', 'Web 端即用', '纯浏览器 SPA 应用\n无需安装，随时随地使用', CLR_BLUE),
    ('🔄', '实时热点追踪', 'Google + 百度双搜索引擎\n实时分析小红书趋势', CLR_GREEN),
    ('🖼️', 'AI 图文一体', '文案 + 配图同步生成\n告别找图烦恼', CLR_PURPLE),
    ('📱', '手机号注册', '一键注册，安全可靠\n手机号唯一，防止滥用', CLR_GOLD),
    ('☁️', '云端部署', 'Render 云平台自动部署\n数据安全，永不丢失', RGBColor(0xFF, 0x85, 0x00)),
]

for i, (icon, title, desc, clr) in enumerate(techs):
    col = i % 3
    row = i // 3
    x = Inches(0.6 + col * 4.1)
    y = Inches(1.5 + row * 3.0)
    add_card(s, x, y, Inches(3.8), Inches(2.6), icon, title, desc.split('\n'), clr)


# ═══════════════════════════════════════
# SLIDE 10 — 总结 & CTA
# ═══════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s, CLR_BG)

# 装饰
circle3 = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10), Inches(5), Inches(5), Inches(5))
circle3.fill.solid(); circle3.fill.fore_color.rgb = CLR_ACCENT
circle3.line.fill.background()
set_shape_alpha(circle3, 12)

add_textbox(s, Inches(1.5), Inches(1.0), Inches(10), Inches(1.2),
            '📝 小红书智能运营台', font_size=48, color=CLR_WHITE, bold=True, align=PP_ALIGN.CENTER)

add_textbox(s, Inches(1.5), Inches(2.3), Inches(10), Inches(0.6),
            '让每一篇笔记都有爆款潜力', font_size=28, color=CLR_ACCENT2, align=PP_ALIGN.CENTER)

# 核心数字
nums = [
    ('7', '技能模板', CLR_ACCENT),
    ('8+', '细分品类', CLR_BLUE),
    ('10x', '效率提升', CLR_GREEN),
    ('3次/日', '免费额度', CLR_GOLD),
]
for i, (num, label, clr) in enumerate(nums):
    x = Inches(1.5 + i * 2.8)
    y = Inches(3.3)
    add_shape_bg(s, x, y, Inches(2.3), Inches(1.5), CLR_CARD_BG, 0.05)
    add_textbox(s, x, y + Inches(0.15), Inches(2.3), Inches(0.8),
                num, font_size=42, color=clr, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(s, x, y + Inches(0.9), Inches(2.3), Inches(0.4),
                label, font_size=16, color=CLR_GRAY, align=PP_ALIGN.CENTER)

# CTA
add_shape_bg(s, Inches(3.5), Inches(5.3), Inches(6.3), Inches(0.9), CLR_ACCENT, 0.15)
add_textbox(s, Inches(3.5), Inches(5.35), Inches(6.3), Inches(0.8),
            '🚀 立即体验：xhs-gemini-proxy.onrender.com',
            font_size=22, color=CLR_WHITE, bold=True, align=PP_ALIGN.CENTER)

add_textbox(s, Inches(1.5), Inches(6.5), Inches(10), Inches(0.5),
            '注册即送免费体验  ·  手机号一键注册  ·  数据安全存储',
            font_size=16, color=CLR_GRAY, align=PP_ALIGN.CENTER)


# ─── 保存 ───
output_path = r'D:\Users\Administrator\Desktop\ls\xhsqg\小红书智能运营台_产品介绍.pptx'
prs.save(output_path)
print(f'✅ PPT 已生成: {output_path}')
